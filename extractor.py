import io
import os
from pathlib import Path
import re
import shutil
from difflib import SequenceMatcher

import pymupdf
from docx import Document
import pytesseract
from PIL import Image


MIN_IMAGE_AREA_RATIO = 0.02
MIN_IMAGE_WIDTH = 100
MIN_IMAGE_HEIGHT = 40


def _tesseract_executable() -> str:
    configured = os.environ.get("TESSERACT_CMD")
    executable = shutil.which(configured or "tesseract")
    if executable is None and configured and Path(configured).is_file():
        executable = configured
    if executable is None:
        raise RuntimeError(
            "Tesseract is required for scanned PDF content. Install Tesseract or "
            "set TESSERACT_CMD to its executable path."
        )
    return executable


def ocr_image(image):
    pytesseract.pytesseract.tesseract_cmd = _tesseract_executable()
    return pytesseract.image_to_string(image).strip()


def _text_blocks(page):
    return [
        block
        for block in page.get_text("blocks", sort=True)
        if len(block) < 7 or block[6] == 0
        if block[4].strip()
    ]


def _extract_page_tables(page, first_table_number, detected_tables):
    tables = []
    table_boxes = []
    header_limit = page.rect.y0 + page.rect.height * 0.2
    footer_limit = page.rect.y0 + page.rect.height * 0.8
    for table in detected_tables:
        box = pymupdf.Rect(table.bbox)
        if box.y1 <= header_limit or box.y0 >= footer_limit:
            continue
        rows = [
            [str(cell).strip() if cell is not None else "" for cell in row]
            for row in table.extract()
        ]
        rows = [row for row in rows if any(cell for cell in row)]
        if rows:
            width = max(len(row) for row in rows)
            rows = [row + [""] * (width - len(row)) for row in rows]
            nonempty_columns = [
                index for index in range(width)
                if any(row[index] for row in rows)
            ]
            rows = [[row[index] for index in nonempty_columns] for row in rows]
        if rows:
            tables.append({"table_id": f"T{first_table_number + len(tables)}", "rows": rows})
            table_boxes.append(box)
    return tables, table_boxes


def _block_is_in_table(block, table_boxes):
    center_x = (block[0] + block[2]) / 2
    center_y = (block[1] + block[3]) / 2
    return any(box.contains(pymupdf.Point(center_x, center_y)) for box in table_boxes)


def _image_boxes(page):
    page_area = max(page.rect.width * page.rect.height, 1)
    boxes = []
    for image in page.get_image_info(xrefs=True):
        box = pymupdf.Rect(image["bbox"]) & page.rect
        if box.is_empty or box.is_infinite:
            continue
        if (
            box.width >= MIN_IMAGE_WIDTH
            and box.height >= MIN_IMAGE_HEIGHT
            and box.width * box.height / page_area >= MIN_IMAGE_AREA_RATIO
        ):
            boxes.append(box)
    return boxes


def _table_paths(page):
    lines = {}
    top = page.rect.y0 + page.rect.height * 0.2
    bottom = page.rect.y0 + page.rect.height * 0.8
    for word in page.get_text("words"):
        x0, y0, x1, _, _, _, _, _ = word
        if top < y0 < bottom:
            lines.setdefault(round(y0 / 4), []).append((x0, x1))
    gap_rows = sum(
        any(right[0] - left[1] > 30 for left, right in zip(sorted(row), sorted(row)[1:]))
        for row in lines.values()
    )
    if gap_rows >= 4:
        return page.get_drawings()
    if gap_rows < 2:
        return None
    paths = page.get_drawings()
    drawing_count = len(paths)
    return paths if drawing_count <= 20 or drawing_count > 40 else None


def _normalise_line(text):
    return re.sub(r"\W+", "", text.casefold(), flags=re.UNICODE)


def _unique_ocr_lines(ocr_text, existing_text):
    existing_lines = [_normalise_line(line) for line in existing_text.splitlines() if line.strip()]
    unique = []
    for line in ocr_text.splitlines():
        normalized = _normalise_line(line)
        if not normalized:
            continue
        duplicate = any(
            normalized == existing
            or (min(len(normalized), len(existing)) >= 12
                and SequenceMatcher(None, normalized, existing).ratio() >= 0.9)
            for existing in existing_lines
        )
        if not duplicate:
            unique.append(line.strip())
    return unique


def _ocr_image_box(page, box):
    pixmap = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), clip=box, alpha=False)
    with Image.open(io.BytesIO(pixmap.tobytes("png"))) as image:
        rgb_image = image.convert("RGB")
        try:
            return ocr_image(rgb_image)
        finally:
            rgb_image.close()


def extract_docx(file_path: str) -> dict:
    doc = Document(file_path)

    paragraphs = [
        p.text.strip()
        for p in doc.paragraphs
        if p.text.strip()
    ]

    tables = []

    for table_no, table in enumerate(doc.tables, start=1):
        rows = []

        for row in table.rows:
            rows.append([
                cell.text.strip()
                for cell in row.cells
            ])

        tables.append({
            "table_id": f"T{table_no}",
            "rows": rows
        })

    return {
        "file_type": "docx",
        "paragraphs": paragraphs,
        "tables": tables
    }


def extract_pdf(file_path: str) -> dict:
    paragraphs = []
    pages = []
    tables = []

    with pymupdf.open(file_path) as pdf:
        for page_no, page in enumerate(pdf, start=1):
            blocks = _text_blocks(page)
            table_paths = _table_paths(page)
            detected_tables = (
                page.find_tables(strategy="lines_strict", paths=table_paths).tables
                if table_paths is not None else []
            )
            page_tables, table_boxes = _extract_page_tables(
                page, len(tables) + 1, detected_tables
            )
            text_blocks = [block for block in blocks if not _block_is_in_table(block, table_boxes)]
            native_paragraphs = [
                "\n".join(line.strip() for line in block[4].splitlines() if line.strip())
                for block in text_blocks
            ]
            native_page_text = "\n".join(native_paragraphs)
            ocr_paragraphs = []
            image_boxes = _image_boxes(page)
            existing_text = native_page_text
            for box in image_boxes:
                ocr_text = _ocr_image_box(page, box)
                unique_lines = _unique_ocr_lines(ocr_text, existing_text)
                ocr_paragraphs.extend(unique_lines)
                existing_text += "\n" + "\n".join(unique_lines)

            page_paragraphs = native_paragraphs + ocr_paragraphs
            paragraphs.extend(page_paragraphs)
            tables.extend(page_tables)
            page_text = "\n".join(page_paragraphs)
            method = "pdf_text+ocr" if ocr_paragraphs and native_paragraphs else "ocr" if ocr_paragraphs else "pdf_text" if native_paragraphs else "none"
            pages.append({
                "page": page_no,
                "extraction_method": method,
                "text": page_text,
                "tables": page_tables,
                "image_regions_ocr": len(image_boxes)
            })

    return {
        "file_type": "pdf",
        "paragraphs": paragraphs,
        "tables": tables,
        "pages": pages
    }


def extract_document(file_path: str) -> dict:
    extension = Path(file_path).suffix.lower()

    if extension == ".docx":
        return extract_docx(file_path)

    if extension == ".pdf":
        return extract_pdf(file_path)

    if extension == ".doc":
        raise ValueError(
            ".doc is not supported directly. Convert it to .docx first."
        )

    raise ValueError(
        f"Unsupported file type: {extension}"
    )